#!/usr/bin/env python3
"""
CPD Breach Report Generator Module for ei_stream_server
=======================================================
Reads 'LM' sheet from CPD Breach / E2E Attribution input file,
filters records where Source DC belongs to the allowed DC Config list,
and generates the formatted output workbook:
  1. 'summary' Sheet: DC breakdown across delay tags ('Customer Attributed',
     'Last Mile delay', 'RTO/IC - NCD') and 'Total CPD Breach', with grand totals.
  2. 'raw' Sheet: Filtered raw records with the 13 required attribution columns.

Uses Single-Pass Zero-Memory Streaming Engine (core.stream_engine):
- O(1) Memory Footprint (< 35MB RAM)
- Direct XML disk streaming for massive datasets
"""

import sys
import gc
import logging
from pathlib import Path
from collections import defaultdict
from typing import Dict, Any, List
import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Ensure server root is in sys.path
SERVER_ROOT = Path(__file__).resolve().parent.parent
if str(SERVER_ROOT) not in sys.path:
    sys.path.insert(0, str(SERVER_ROOT))

try:
    from config.dc_config import ALLOWED_DCS_SET_LOWER
except ImportError:
    ALLOWED_DCS_SET_LOWER = {'alg', 'ayp', 'deo', 'jhs', 'jnp', 'knp', 'mau', 'mrz', 'mth', 'mzn', 'rbr', 'spr', 'vns', 'all'}

from core.stream_engine import (
    XmlSheetWriter,
    assemble_stream_workbook,
    open_stream_reader,
    get_sheet_names,
    ColumnFinder
)

log = logging.getLogger("ei_stream_server.cpd_breach_report")

RAW_COLUMNS = [
    'tracking_number',
    'order_created_date',
    'customer_promise_date',
    'New_Final_delay_tag',
    'rto_ic_flag',
    'Latest Status',
    'Current Location',
    'Latest Update Time',
    'Cs Notes',
    'No. of Attempts',
    'DC Code',
    'Source DC',
    'Region'
]

DELAY_TAGS = ['Customer Attributed', 'Last Mile delay', 'RTO/IC - NCD']


def generate_cpd_breach_report(input_file: Path, output_file: Path):
    """
    Zero-memory streaming generator for CPD Breach Report.
    Reads input file, filters for DC Config hubs, streams 'raw' rows to disk,
    and stitches with OpenPyXL styled 'summary' sheet.
    """
    input_path = Path(input_file)
    output_path = Path(output_file)
    log.info(f"Loading input workbook for CPD Breach Report (Single-Pass Stream): {input_path.name}")

    sheet_names = get_sheet_names(input_path)
    sheet_map = {name.lower().strip(): name for name in sheet_names}

    target_sheet = None
    for candidate in ['lm', 'lm_sheet', 'raw', 'raw_data', 'data']:
        if candidate in sheet_map:
            target_sheet = sheet_map[candidate]
            break
    if not target_sheet and sheet_names:
        target_sheet = sheet_names[0]

    sum_stats: Dict[str, Dict[str, int]] = defaultdict(
        lambda: {tag: 0 for tag in DELAY_TAGS} | {'Total CPD Breach': 0}
    )
    raw_writer = XmlSheetWriter("raw", RAW_COLUMNS)

    with open_stream_reader(input_path, sheet_name=target_sheet) as (raw_headers, row_iter):
        if not raw_headers:
            raise ValueError(f"Target sheet '{target_sheet}' in {input_path.name} is empty or missing headers.")

        cf = ColumnFinder(raw_headers, {
            'tracking_number': ['tracking_number', 'tracking number', 'tracking_no', 'waybill'],
            'order_created_date': ['order_created_date', 'order created date', 'created_date', 'order_date'],
            'customer_promise_date': ['customer_promise_date', 'customer promise date', 'cpd', 'promise_date'],
            'New_Final_delay_tag': ['new_final_delay_tag', 'delay tag', 'delay_tag', 'final_delay_tag'],
            'rto_ic_flag': ['rto_ic_flag', 'rto ic flag', 'rto_flag'],
            'Latest Status': ['latest status', 'latest_status', 'status'],
            'Current Location': ['current location', 'current_location', 'location'],
            'Latest Update Time': ['latest update time', 'latest_update_time', 'update_time'],
            'Cs Notes': ['cs notes', 'cs_notes', 'notes', 'cs note'],
            'No. of Attempts': ['no. of attempts', 'no of attempts', 'attempts', 'no_of_attempts'],
            'DC Code': ['dc code', 'dc_code', 'dccode'],
            'Source DC': ['source dc', 'sourcedc', 'source_dc', 'dc', 'hub'],
            'Region': ['region', 'zone']
        })

        col_indices = [cf.get(col, -1) for col in RAW_COLUMNS]
        source_dc_idx = cf.get('Source DC', -1)
        delay_tag_idx = cf.get('New_Final_delay_tag', -1)

        if source_dc_idx == -1:
            # Fallback to checking raw header exact match
            lower_headers = [str(h).strip().lower() for h in raw_headers]
            if 'source dc' in lower_headers:
                source_dc_idx = lower_headers.index('source dc')

        filtered_count = 0

        with raw_writer:
            for row in row_iter:
                if not row or len(row) <= source_dc_idx or source_dc_idx < 0:
                    continue

                raw_dc = row[source_dc_idx]
                if raw_dc is None:
                    continue

                dc_clean = str(raw_dc).strip().lower()
                if dc_clean in ALLOWED_DCS_SET_LOWER:
                    # Map row to the 13 required columns
                    row_vals: List[Any] = []
                    for c_idx in col_indices:
                        if c_idx >= 0 and c_idx < len(row):
                            val = row[c_idx]
                            row_vals.append("" if val is None else val)
                        else:
                            row_vals.append("")

                    raw_writer.write_row(row_vals)
                    filtered_count += 1

                    # Aggregate summary metrics
                    dc_upper = str(raw_dc).strip().upper()
                    tag_val = str(row[delay_tag_idx]).strip() if delay_tag_idx >= 0 and delay_tag_idx < len(row) and row[delay_tag_idx] is not None else ""

                    if tag_val in sum_stats[dc_upper]:
                        sum_stats[dc_upper][tag_val] += 1
                    sum_stats[dc_upper]['Total CPD Breach'] += 1

    log.info(f"Filtered {filtered_count} records for CPD Breach Report across {len(sum_stats)} allowed DCs.")

    # -------------------------------------------------------------------------
    # Build OpenPyXL Micro-Workbook for 'summary' sheet (< 2 MB RAM)
    # -------------------------------------------------------------------------
    wb = Workbook()
    ws_sum = wb.active
    ws_sum.title = "summary"
    ws_sum.sheet_view.showGridLines = True

    # Empty placeholder sheet for raw data (replaced during assemble_stream_workbook)
    wb.create_sheet(title="raw")

    # Styling Palette
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")  # Navy Blue
    total_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")   # Soft Blue
    alt_fill = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")     # Subtle zebra

    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    font_bold = Font(name="Calibri", size=11, bold=True)
    font_regular = Font(name="Calibri", size=11)

    thin_border = Side(border_style="thin", color="D9D9D9")
    double_border = Side(border_style="double", color="000000")
    top_thin_border = Side(border_style="thin", color="000000")

    cell_border = Border(left=thin_border, right=thin_border, top=thin_border, bottom=thin_border)
    total_border = Border(left=thin_border, right=thin_border, top=top_thin_border, bottom=double_border)

    sum_headers = ['Source DC', 'Customer Attributed', 'Last Mile delay', 'RTO/IC - NCD', 'Total CPD Breach']
    ws_sum.row_dimensions[1].height = 25

    for col_idx, h_text in enumerate(sum_headers, start=1):
        cell = ws_sum.cell(row=1, column=col_idx, value=h_text)
        cell.font = font_header
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = cell_border

    sorted_dcs = sorted(sum_stats.keys())
    row_idx = 2

    grand_totals = {tag: 0 for tag in DELAY_TAGS}
    grand_totals['Total CPD Breach'] = 0

    for sdc in sorted_dcs:
        stats = sum_stats[sdc]
        c_att = stats['Customer Attributed']
        lm_del = stats['Last Mile delay']
        rto_ncd = stats['RTO/IC - NCD']
        tot = stats['Total CPD Breach']

        grand_totals['Customer Attributed'] += c_att
        grand_totals['Last Mile delay'] += lm_del
        grand_totals['RTO/IC - NCD'] += rto_ncd
        grand_totals['Total CPD Breach'] += tot

        ws_sum.row_dimensions[row_idx].height = 20
        use_alt = (row_idx % 2 == 1)

        c1 = ws_sum.cell(row=row_idx, column=1, value=sdc)
        c1.font = font_bold
        c1.alignment = Alignment(horizontal="center", vertical="center")
        c1.border = cell_border
        if use_alt:
            c1.fill = alt_fill

        data_values = [
            c_att if c_att > 0 else "",
            lm_del if lm_del > 0 else "",
            rto_ncd if rto_ncd > 0 else "",
            tot if tot > 0 else ""
        ]

        for c_i, val in enumerate(data_values, start=2):
            cell = ws_sum.cell(row=row_idx, column=c_i, value=val)
            cell.font = font_regular
            cell.alignment = Alignment(horizontal="right", vertical="center")
            cell.border = cell_border
            if val != "":
                cell.number_format = "#,##0"
            if use_alt:
                cell.fill = alt_fill

        row_idx += 1

    # Total Result Row
    ws_sum.row_dimensions[row_idx].height = 22
    t_cell1 = ws_sum.cell(row=row_idx, column=1, value="Total Result")
    t_cell1.font = font_bold
    t_cell1.alignment = Alignment(horizontal="center", vertical="center")
    t_cell1.fill = total_fill
    t_cell1.border = total_border

    tot_values = [
        grand_totals['Customer Attributed'],
        grand_totals['Last Mile delay'],
        grand_totals['RTO/IC - NCD'],
        grand_totals['Total CPD Breach']
    ]

    for c_i, val in enumerate(tot_values, start=2):
        c = ws_sum.cell(row=row_idx, column=c_i, value=val)
        c.font = font_bold
        c.alignment = Alignment(horizontal="right", vertical="center")
        c.fill = total_fill
        c.border = total_border
        c.number_format = "#,##0"

    # Column widths auto-fit
    for col in ws_sum.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if isinstance(cell.value, (int, float)):
                val_str = f"{cell.value:,.0f}"
            max_len = max(max_len, len(val_str))
        ws_sum.column_dimensions[col_letter].width = max(max_len + 4, 16)

    # -------------------------------------------------------------------------
    # Assemble Output (Replaces 'raw' placeholder with streamed disk XML)
    # -------------------------------------------------------------------------
    assemble_stream_workbook(wb, [raw_writer], output_path)
    raw_writer.cleanup()
    del wb
    gc.collect()

    log.info(f"Successfully generated CPD Breach Report: {output_path.name}")
